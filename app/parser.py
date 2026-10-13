from __future__ import annotations

import colorsys
from dataclasses import dataclass

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.shapes.group import GroupShape

from app.config import MONOSPACE_FONTS
from app.models import (
    ArchiveFacts,
    ElementModel,
    LineModel,
    ParagraphModel,
    PresentationModel,
    RunModel,
    SlideModel,
)

TITLE_KINDS = {"title", "center_title", "vertical_title"}
DECORATION_KINDS = {"footer", "slide_number", "date"}

PLACEHOLDER_KINDS = {
    PP_PLACEHOLDER.TITLE: "title",
    PP_PLACEHOLDER.CENTER_TITLE: "center_title",
    PP_PLACEHOLDER.VERTICAL_TITLE: "vertical_title",
    PP_PLACEHOLDER.SUBTITLE: "body",
    PP_PLACEHOLDER.BODY: "body",
    PP_PLACEHOLDER.VERTICAL_BODY: "body",
    PP_PLACEHOLDER.OBJECT: "body",
    PP_PLACEHOLDER.VERTICAL_OBJECT: "body",
    PP_PLACEHOLDER.FOOTER: "footer",
    PP_PLACEHOLDER.SLIDE_NUMBER: "slide_number",
    PP_PLACEHOLDER.DATE: "date",
    PP_PLACEHOLDER.HEADER: "other",
}

SCHEME_ALIASES = {
    "bg1": "lt1",
    "tx1": "dk1",
    "bg2": "lt2",
    "tx2": "dk2",
}


@dataclass
class Transform:
    ox: float
    oy: float
    sx: float
    sy: float

    def apply(self, x: int, y: int, w: int, h: int) -> tuple[int, int, int, int]:
        return (
            int(round(self.ox + x * self.sx)),
            int(round(self.oy + y * self.sy)),
            int(round(w * self.sx)),
            int(round(h * self.sy)),
        )


IDENTITY = Transform(0, 0, 1, 1)


@dataclass
class Theme:
    colors: dict[str, tuple[int, int, int]]
    color_map: dict[str, str]
    major_latin: str | None
    minor_latin: str | None

    def resolve_scheme(self, name: str) -> tuple[int, int, int] | None:
        mapped = self.color_map.get(name, SCHEME_ALIASES.get(name, name))
        mapped = self.color_map.get(mapped, mapped)
        return self.colors.get(mapped) or self.colors.get(name)

    def resolve_font(self, typeface: str | None) -> str | None:
        if not typeface:
            return None
        if typeface.startswith("+mj"):
            return self.major_latin
        if typeface.startswith("+mn"):
            return self.minor_latin
        return typeface


def parse_pptx(presentation, archive_names: list[str]) -> PresentationModel:
    width = int(presentation.slide_width)
    height = int(presentation.slide_height)
    themes = _theme_cache(presentation)
    slides: list[SlideModel] = []
    for number, slide in enumerate(presentation.slides, start=1):
        theme = themes.get(id(slide.slide_layout.slide_master), Theme({}, {}, None, None))
        elements: list[ElementModel] = []
        counter = _Counter()
        _walk_shapes(
            slide.shapes,
            IDENTITY,
            slide,
            presentation,
            theme,
            number,
            counter,
            elements,
        )
        has_video = any(element.type == "video" for element in elements) or _slide_has_video(slide)
        has_title = any(
            element.placeholder_kind in TITLE_KINDS for element in elements
        )
        slides.append(
            SlideModel(
                number=number,
                width_emu=width,
                height_emu=height,
                elements=elements,
                has_video=has_video,
                has_title_placeholder=has_title,
            )
        )
    has_fonts = any(name.startswith("ppt/fonts/") for name in archive_names)
    return PresentationModel(
        width_emu=width,
        height_emu=height,
        slides=slides,
        archive=ArchiveFacts(has_fonts_dir=has_fonts),
    )


class _Counter:
    def __init__(self) -> None:
        self.value = 0

    def next(self, slide_number: int) -> str:
        self.value += 1
        return f"s{slide_number}-e{self.value}"


def _walk_shapes(shapes, transform, slide, presentation, theme, number, counter, out) -> None:
    for shape in shapes:
        if isinstance(shape, GroupShape) or shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _append_group(shape, transform, slide, presentation, theme, number, counter, out)
            continue
        _append_leaf(shape, transform, slide, presentation, theme, number, counter, out)


def _append_group(shape, transform, slide, presentation, theme, number, counter, out) -> None:
    gx, gy, gw, gh, chx, chy, chw, chh = _group_box(shape)
    ax, ay, aw, ah = transform.apply(gx, gy, gw, gh)
    out.append(
        ElementModel(
            id=counter.next(number),
            type="group",
            x=ax,
            y=ay,
            w=aw,
            h=ah,
            text="",
            is_title=False,
            placeholder_kind=None,
            is_decoration=False,
            paragraphs=[],
            fill_type=_fill(shape._element, theme)[0],
            fill_rgb=_fill(shape._element, theme)[1],
            fill_transparent=_fill(shape._element, theme)[2],
        )
    )
    sx = (gw / chw) if chw else 1
    sy = (gh / chh) if chh else 1
    child_tf = Transform(
        ox=transform.ox + transform.sx * gx - transform.sx * sx * chx,
        oy=transform.oy + transform.sy * gy - transform.sy * sy * chy,
        sx=transform.sx * sx,
        sy=transform.sy * sy,
    )
    _walk_shapes(shape.shapes, child_tf, slide, presentation, theme, number, counter, out)


def _append_leaf(shape, transform, slide, presentation, theme, number, counter, out) -> None:
    x, y, w, h = _shape_box(shape)
    ax, ay, aw, ah = transform.apply(x, y, w, h)
    kind = _placeholder_kind(shape)
    is_title = kind in TITLE_KINDS
    is_decoration = kind in DECORATION_KINDS
    element_type = _element_type(shape, kind)
    paragraphs = _paragraphs(shape, slide, presentation, theme, kind)
    text = "\n".join(paragraph.text for paragraph in paragraphs if paragraph.text).strip()
    fill_type, fill_rgb, fill_transparent = _fill(shape._element, theme)
    if element_type == "picture" and _element_has_video(shape._element):
        element_type = "video"
    out.append(
        ElementModel(
            id=counter.next(number),
            type=element_type,
            x=ax,
            y=ay,
            w=aw,
            h=ah,
            text=text,
            is_title=is_title,
            placeholder_kind=kind,
            is_decoration=is_decoration,
            paragraphs=paragraphs,
            fill_type=fill_type,
            fill_rgb=fill_rgb,
            fill_transparent=fill_transparent,
        )
    )


def _element_type(shape, kind: str | None) -> str:
    if _element_has_video(shape._element) or shape.shape_type == MSO_SHAPE_TYPE.MEDIA:
        if _is_audio(shape):
            return "shape"
        return "video"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if shape.shape_type == MSO_SHAPE_TYPE.TABLE or getattr(shape, "has_table", False):
        return "table"
    if kind in DECORATION_KINDS:
        return "placeholder"
    if kind is not None or shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
        return "text"
    if getattr(shape, "has_text_frame", False):
        try:
            if shape.text_frame.text.strip():
                return "text"
        except Exception:
            pass
    return "shape"


def _is_audio(shape) -> bool:
    try:
        for rel in shape.part.rels.values():
            if "audio" in (rel.reltype or ""):
                return True
    except Exception:
        return False
    return False


def _element_has_video(element) -> bool:
    return bool(element.findall(".//" + qn("a:videoFile")) or element.findall(".//" + qn("p:video")))


def _slide_has_video(slide) -> bool:
    try:
        for rel in slide.part.rels.values():
            reltype = rel.reltype or ""
            if reltype.endswith("/video"):
                return True
    except Exception:
        return False
    return False


def _placeholder_kind(shape) -> str | None:
    if not getattr(shape, "is_placeholder", False):
        return None
    try:
        ph_type = shape.placeholder_format.type
    except Exception:
        return "other"
    if ph_type is None:
        return "other"
    return PLACEHOLDER_KINDS.get(ph_type, "other")


def _shape_box(shape) -> tuple[int, int, int, int]:
    box = _xfrm_box(shape._element)
    if box is not None:
        return box
    return (
        int(shape.left or 0),
        int(shape.top or 0),
        int(shape.width or 0),
        int(shape.height or 0),
    )


def _group_box(shape) -> tuple[int, int, int, int, int, int, int, int]:
    xfrm = _direct_xfrm(shape._element)
    if xfrm is None:
        x, y, w, h = _shape_box(shape)
        return x, y, w, h, 0, 0, w or 1, h or 1
    x, y, w, h = _off_ext(xfrm)
    ch_off = xfrm.find(qn("a:chOff"))
    ch_ext = xfrm.find(qn("a:chExt"))
    chx = int(ch_off.get("x")) if ch_off is not None and ch_off.get("x") else 0
    chy = int(ch_off.get("y")) if ch_off is not None and ch_off.get("y") else 0
    chw = int(ch_ext.get("cx")) if ch_ext is not None and ch_ext.get("cx") else w
    chh = int(ch_ext.get("cy")) if ch_ext is not None and ch_ext.get("cy") else h
    return x, y, w, h, chx, chy, chw or 1, chh or 1


def _xfrm_box(element) -> tuple[int, int, int, int] | None:
    xfrm = _direct_xfrm(element)
    if xfrm is None:
        return None
    return _off_ext(xfrm)


def _direct_xfrm(element):
    for tag in ("p:spPr", "p:grpSpPr"):
        node = element.find(qn(tag))
        if node is not None:
            xfrm = node.find(qn("a:xfrm"))
            if xfrm is not None:
                return xfrm
    return element.find(qn("p:xfrm"))


def _off_ext(xfrm) -> tuple[int, int, int, int]:
    off = xfrm.find(qn("a:off"))
    ext = xfrm.find(qn("a:ext"))
    x = int(off.get("x")) if off is not None and off.get("x") else 0
    y = int(off.get("y")) if off is not None and off.get("y") else 0
    w = int(ext.get("cx")) if ext is not None and ext.get("cx") else 0
    h = int(ext.get("cy")) if ext is not None and ext.get("cy") else 0
    return x, y, w, h


def _paragraphs(shape, slide, presentation, theme, kind: str | None) -> list[ParagraphModel]:
    frames = []
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                frames.append(cell.text_frame)
    elif getattr(shape, "has_text_frame", False):
        frames.append(shape.text_frame)
    paragraphs: list[ParagraphModel] = []
    for frame in frames:
        if frame is None:
            continue
        for paragraph in frame.paragraphs:
            parsed = _paragraph(paragraph._p, shape, slide, presentation, theme, kind)
            if parsed is not None:
                paragraphs.append(parsed)
    return paragraphs


def _paragraph(p_elm, shape, slide, presentation, theme, kind) -> ParagraphModel | None:
    level = 0
    p_pr = p_elm.find(qn("a:pPr"))
    if p_pr is not None and p_pr.get("lvl"):
        level = int(p_pr.get("lvl"))
    bullet, numbered = _list_flags(p_elm, shape, level)
    sources_without_run = _rpr_sources(None, p_elm, shape, slide, presentation, kind, level)
    lines: list[LineModel] = []
    current: list[RunModel] = []

    def flush(explicit: bool) -> None:
        nonlocal current
        if current or explicit:
            lines.append(LineModel(runs=current, explicit_break_after=explicit))
            current = []

    for child in p_elm:
        tag = child.tag
        if tag == qn("a:r") or tag == qn("a:fld"):
            run = _run(child, p_elm, shape, slide, presentation, theme, kind, level, sources_without_run)
            if run is not None and run.text:
                current.append(run)
        elif tag == qn("a:br"):
            flush(True)
    if current:
        flush(False)
    if not lines:
        return ParagraphModel(level=level, bullet=bullet, numbered=numbered, lines=[])
    if lines and lines[-1].explicit_break_after and not lines[-1].runs:
        pass
    return ParagraphModel(level=level, bullet=bullet, numbered=numbered, lines=lines)


def _run(r_elm, p_elm, shape, slide, presentation, theme, kind, level, inherited_sources) -> RunModel | None:
    texts = r_elm.findall(qn("a:t"))
    text = "".join(node.text or "" for node in texts)
    if text == "":
        return None
    r_pr = r_elm.find(qn("a:rPr"))
    sources = ([r_pr] if r_pr is not None else []) + inherited_sources
    font = _first(sources, lambda item: _font_from_rpr(item, theme))
    size = _first(sources, _size_from_rpr)
    color, text_transparent = _first_color(sources, theme)
    superscript = _first(sources, _superscript_from_rpr)
    font_key = font.casefold() if font else None
    return RunModel(
        text=text,
        font_name=font,
        font_size_pt=size,
        color_rgb=color,
        superscript=bool(superscript),
        monospace=font_key in MONOSPACE_FONTS if font_key else False,
        text_transparent=text_transparent,
    )


def _rpr_sources(r_pr, p_elm, shape, slide, presentation, kind, level) -> list:
    sources = []
    p_pr = p_elm.find(qn("a:pPr")) if p_elm is not None else None
    if p_pr is not None:
        def_rpr = p_pr.find(qn("a:defRPr"))
        if def_rpr is not None:
            sources.append(def_rpr)
    tx_body = shape._element.find(qn("p:txBody"))
    if tx_body is not None:
        lst = tx_body.find(qn("a:lstStyle"))
        found = _lvl_def_rpr(lst, level)
        if found is not None:
            sources.append(found)
    layout_shape = _layout_placeholder(shape, slide)
    if layout_shape is not None:
        sources.extend(_shape_def_rprs(layout_shape._element, level))
    master = slide.slide_layout.slide_master
    style_tag = _master_style_tag(kind)
    tx_styles = master.element.find(qn("p:txStyles"))
    if tx_styles is not None:
        style = tx_styles.find(qn(style_tag))
        found = _lvl_def_rpr(style, level)
        if found is not None:
            sources.append(found)
    default_text = presentation.element.find(qn("p:defaultTextStyle"))
    found = _lvl_def_rpr(default_text, level)
    if found is not None:
        sources.append(found)
    return sources


def _shape_def_rprs(element, level) -> list:
    found = []
    tx_body = element.find(qn("p:txBody"))
    if tx_body is None:
        return found
    lst = tx_body.find(qn("a:lstStyle"))
    lvl = _lvl_def_rpr(lst, level)
    if lvl is not None:
        found.append(lvl)
    for paragraph in tx_body.findall(qn("a:p")):
        p_pr = paragraph.find(qn("a:pPr"))
        if p_pr is not None and p_pr.find(qn("a:defRPr")) is not None:
            found.append(p_pr.find(qn("a:defRPr")))
        for run in paragraph.findall(qn("a:r")):
            r_pr = run.find(qn("a:rPr"))
            if r_pr is not None:
                found.append(r_pr)
    return found


def _lvl_def_rpr(parent, level: int):
    if parent is None:
        return None
    lvl = parent.find(qn(f"a:lvl{level + 1}pPr"))
    if lvl is None:
        return None
    return lvl.find(qn("a:defRPr"))


def _layout_placeholder(shape, slide):
    if not getattr(shape, "is_placeholder", False):
        return None
    try:
        idx = shape.placeholder_format.idx
    except Exception:
        return None
    for candidate in slide.slide_layout.shapes:
        if getattr(candidate, "is_placeholder", False) and candidate.placeholder_format.idx == idx:
            return candidate
    return None


def _master_style_tag(kind: str | None) -> str:
    if kind in TITLE_KINDS:
        return "p:titleStyle"
    if kind == "body":
        return "p:bodyStyle"
    return "p:otherStyle"


def _list_flags(p_elm, shape, level: int) -> tuple[bool, bool]:
    p_pr = p_elm.find(qn("a:pPr"))
    if p_pr is not None and p_pr.find(qn("a:buNone")) is not None:
        return False, False
    numbered = _has_auto_num(p_pr)
    bullet = _has_bullet(p_pr)
    if bullet or numbered:
        return True, numbered
    tx_body = shape._element.find(qn("p:txBody"))
    lst = tx_body.find(qn("a:lstStyle")) if tx_body is not None else None
    lvl = lst.find(qn(f"a:lvl{level + 1}pPr")) if lst is not None else None
    if lvl is not None and lvl.find(qn("a:buNone")) is None:
        return _has_bullet(lvl) or _has_auto_num(lvl), _has_auto_num(lvl)
    return False, False


def _has_bullet(node) -> bool:
    if node is None:
        return False
    return node.find(qn("a:buChar")) is not None or node.find(qn("a:buBlip")) is not None


def _has_auto_num(node) -> bool:
    return node is not None and node.find(qn("a:buAutoNum")) is not None


def _first(sources, extractor):
    for source in sources:
        if source is None:
            continue
        value = extractor(source)
        if value is not None:
            return value
    return None


def _first_color(sources, theme: Theme) -> tuple[tuple[int, int, int] | None, bool]:
    for source in sources:
        if source is None:
            continue
        parsed = _color_from_parent(source, theme)
        if parsed is not None:
            return parsed
    return None, False


def _size_from_rpr(r_pr) -> float | None:
    raw = r_pr.get("sz")
    if raw is None:
        return None
    try:
        return int(raw) / 100
    except ValueError:
        return None


def _font_from_rpr(r_pr, theme: Theme) -> str | None:
    for tag in ("a:latin", "a:cs", "a:ea"):
        node = r_pr.find(qn(tag))
        if node is None:
            continue
        typeface = node.get("typeface")
        resolved = theme.resolve_font(typeface)
        if resolved:
            return resolved
    return None


def _superscript_from_rpr(r_pr) -> bool | None:
    raw = r_pr.get("baseline")
    if raw is None:
        return None
    try:
        return int(raw) > 0
    except ValueError:
        return None


def _fill(element, theme: Theme) -> tuple[str | None, tuple[int, int, int] | None, bool]:
    sp_pr = element.find(qn("p:spPr"))
    if sp_pr is None:
        sp_pr = element.find(qn("p:grpSpPr"))
    if sp_pr is None:
        return None, None, False
    if sp_pr.find(qn("a:noFill")) is not None:
        return "none", None, False
    if sp_pr.find(qn("a:gradFill")) is not None:
        return "gradient", None, False
    if sp_pr.find(qn("a:blipFill")) is not None or sp_pr.find(qn("a:pattFill")) is not None:
        return "picture", None, False
    solid = sp_pr.find(qn("a:solidFill"))
    if solid is not None:
        parsed = _color_from_parent(solid, theme)
        if parsed is None:
            return "solid", None, False
        rgb, transparent = parsed
        return "solid", rgb, transparent
    return None, None, False


def _color_from_parent(parent, theme: Theme) -> tuple[tuple[int, int, int], bool] | None:
    for child in parent:
        tag = child.tag
        if tag not in {qn("a:solidFill"), qn("a:srgbClr"), qn("a:schemeClr"), qn("a:sysClr"), qn("a:prstClr")}:
            continue
        if tag == qn("a:solidFill"):
            return _color_from_parent(child, theme)
        return _color_node(child, theme)
    return None


def _color_node(node, theme: Theme) -> tuple[tuple[int, int, int], bool] | None:
    tag = node.tag
    rgb: tuple[int, int, int] | None = None
    if tag == qn("a:srgbClr"):
        rgb = _hex_rgb(node.get("val"))
    elif tag == qn("a:sysClr"):
        rgb = _hex_rgb(node.get("lastClr"))
    elif tag == qn("a:schemeClr"):
        name = node.get("val")
        if name:
            rgb = theme.resolve_scheme(name)
    elif tag == qn("a:prstClr"):
        rgb = {"black": (0, 0, 0), "white": (255, 255, 255)}.get(node.get("val") or "")
    if rgb is None:
        return None
    transparent = False
    for child in node:
        ctag = child.tag
        val = _int_attr(child, "val")
        if ctag == qn("a:alpha") and val is not None and val < 100_000:
            transparent = True
        elif ctag == qn("a:lumMod") and val is not None:
            rgb = _apply_lum(rgb, val, None)
        elif ctag == qn("a:lumOff") and val is not None:
            rgb = _apply_lum(rgb, None, val)
        elif ctag == qn("a:shade") and val is not None:
            rgb = tuple(max(0, min(255, round(channel * val / 100_000))) for channel in rgb)  # type: ignore[assignment]
        elif ctag == qn("a:tint") and val is not None:
            factor = val / 100_000
            rgb = tuple(
                max(0, min(255, round(channel + (255 - channel) * (1 - factor))))
                for channel in rgb
            )  # type: ignore[assignment]
    return rgb, transparent


def _apply_lum(rgb: tuple[int, int, int], lum_mod: int | None, lum_off: int | None) -> tuple[int, int, int]:
    red, green, blue = (channel / 255 for channel in rgb)
    hue, light, sat = colorsys.rgb_to_hls(red, green, blue)
    if lum_mod is not None:
        light *= lum_mod / 100_000
    if lum_off is not None:
        light += lum_off / 100_000
    light = min(1.0, max(0.0, light))
    red, green, blue = colorsys.hls_to_rgb(hue, light, sat)
    return (round(red * 255), round(green * 255), round(blue * 255))


def _hex_rgb(value: str | None) -> tuple[int, int, int] | None:
    if not value or len(value) < 6:
        return None
    try:
        return (int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
    except ValueError:
        return None


def _int_attr(node, name: str) -> int | None:
    raw = node.get(name)
    if raw is None:
        return None
    try:
        return int(raw)
    except ValueError:
        return None


def _theme_cache(presentation) -> dict[int, Theme]:
    cache: dict[int, Theme] = {}
    for master in presentation.slide_masters:
        if id(master) in cache:
            continue
        colors: dict[str, tuple[int, int, int]] = {}
        major = None
        minor = None
        theme_elm = _theme_element(master)
        if theme_elm is not None:
            scheme = theme_elm.find(".//" + qn("a:clrScheme"))
            if scheme is not None:
                for child in scheme:
                    name = etree.QName(child).localname
                    parsed = _color_from_parent(child, Theme({}, {}, None, None))
                    if parsed is not None:
                        colors[name] = parsed[0]
            font_scheme = theme_elm.find(".//" + qn("a:fontScheme"))
            if font_scheme is not None:
                major_latin = font_scheme.find("./" + qn("a:majorFont") + "/" + qn("a:latin"))
                minor_latin = font_scheme.find("./" + qn("a:minorFont") + "/" + qn("a:latin"))
                if major_latin is not None:
                    major = major_latin.get("typeface")
                if minor_latin is not None:
                    minor = minor_latin.get("typeface")
        color_map = dict(SCHEME_ALIASES)
        clr_map = master.element.find(qn("p:clrMap"))
        if clr_map is not None:
            color_map.update(dict(clr_map.attrib))
        cache[id(master)] = Theme(colors, color_map, major, minor)
    return cache


def _theme_element(master):
    try:
        for rel in master.part.rels.values():
            if not (rel.reltype or "").endswith("/theme"):
                continue
            part = rel.target_part
            element = getattr(part, "element", None)
            if element is None:
                element = getattr(part, "_element", None)
            if element is None and getattr(part, "blob", None):
                element = etree.fromstring(part.blob)
            return element
    except Exception:
        return None
    return None
