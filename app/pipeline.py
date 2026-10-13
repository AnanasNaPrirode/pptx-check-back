from __future__ import annotations

import json
from io import BytesIO
from pathlib import PurePath

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from app.checks import Check, catalog, run_check
from app.errors import ApiError
from app.logging_setup import log
from app.models import PresentationModel
from app.parser import parse_pptx
from app.validate import validate_upload


def parse_enabled_checks(raw: str | None) -> list[Check]:
    items = catalog()
    if raw is None or raw.strip() == "":
        return [item for item in items if item.default_enabled]
    try:
        parsed = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ApiError(422, "UNKNOWN_CHECK", "Некорректный список проверок.") from exc
    if not isinstance(parsed, list) or not all(isinstance(item, str) for item in parsed):
        raise ApiError(422, "UNKNOWN_CHECK", "Некорректный список проверок.")
    known = {item.id for item in items}
    unknown = [item for item in parsed if item not in known]
    if unknown:
        raise ApiError(422, "UNKNOWN_CHECK", f"Неизвестная проверка: {unknown[0]}.")
    selected = set(parsed)
    return [item for item in items if item.id in selected]


def build_report(filename: str | None, content_type: str | None, data: bytes, enabled_raw: str | None) -> dict:
    checks = parse_enabled_checks(enabled_raw)
    names = validate_upload(filename, content_type, data)
    try:
        presentation = Presentation(BytesIO(data))
    except (PackageNotFoundError, ValueError, KeyError, OSError) as exc:
        raise ApiError(400, "INVALID_FILE", "Файл повреждён или не является презентацией .pptx.") from exc
    except Exception as exc:
        raise ApiError(400, "INVALID_FILE", "Файл повреждён или не является презентацией .pptx.") from exc
    model = parse_pptx(presentation, names)
    _apply(model, checks)
    safe_name = PurePath(filename or "presentation.pptx").name
    return {
        "filename": safe_name,
        "slides_count": len(model.slides),
        "applied_checks": [item.id for item in checks],
        "presentation_issues": [issue.presentation_public() for issue in model.presentation_issues],
        "slides": [slide.public() for slide in model.slides],
    }


def _apply(model: PresentationModel, checks: list[Check]) -> None:
    for check in checks:
        try:
            run_check(check, model)
        except Exception as exc:
            log.error(
                "check_failed",
                extra={"check_id": check.id, "error_code": type(exc).__name__},
            )
