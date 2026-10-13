import json

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from app.pipeline import build_report
from tests.pptx_factory import MIN_PNG, SLIDE_H, add_text, blank, presentation, save, set_bullet


def issues_for(data: bytes, check_id: str) -> list[dict]:
    report = build_report("t.pptx", None, data, json.dumps([check_id]))
    found = list(report["presentation_issues"])
    for slide in report["slides"]:
        found.extend(slide["issues"])
    return found


def assert_clean(builder, check_id: str) -> None:
    prs = presentation()
    builder(prs, False)
    assert issues_for(save(prs), check_id) == []


def assert_flagged(builder, check_id: str) -> dict:
    prs = presentation()
    builder(prs, True)
    found = issues_for(save(prs), check_id)
    assert len(found) == 1
    assert found[0]["check_id"] == check_id
    assert found[0]["message"]
    return found[0]


def _aspect(prs, bad: bool) -> None:
    if bad:
        prs.slide_width = Emu(9_144_000)
    blank(prs)


def _fonts(prs, bad: bool) -> None:
    slide = blank(prs)
    add_text(slide, "Текст", 400_000, 400_000, 4_000_000, 800_000, font="RareFontXYZ" if bad else "Calibri")


def _duplicates(prs, bad: bool) -> None:
    first = blank(prs)
    add_text(first, "Одинаково", 400_000, 400_000, 4_000_000, 800_000)
    second = blank(prs)
    add_text(second, "Одинаково" if bad else "Иначе", 400_000, 400_000, 4_000_000, 800_000)


def _missing_title(prs, bad: bool) -> None:
    if bad:
        slide = blank(prs)
        add_text(slide, "Без плейсхолдера", 400_000, 400_000, 6_000_000, 800_000)
        return
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Тема"


def _title_dot(prs, bad: bool) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Тема." if bad else "Тема"


def _video(prs, bad: bool) -> None:
    slide = blank(prs)
    if not bad:
        add_text(slide, "Без видео", 400_000, 400_000, 4_000_000, 800_000)
        return
    picture = slide.shapes.add_picture(BytesIO_PNG(), Emu(400_000), Emu(400_000), Emu(800_000), Emu(800_000))
    nv_pr = picture._element.find(".//" + qn("p:nvPr"))
    etree.SubElement(nv_pr, qn("a:videoFile")).set(qn("r:link"), "rId9")


def BytesIO_PNG():
    from io import BytesIO

    return BytesIO(MIN_PNG)


def _quotes(prs, bad: bool) -> None:
    slide = blank(prs)
    add_text(slide, 'Он сказал "привет"' if bad else "Он сказал «привет»", 400_000, 400_000, 6_000_000, 800_000)


def _dash(prs, bad: bool) -> None:
    slide = blank(prs)
    add_text(slide, "это - пример" if bad else "это — пример", 400_000, 400_000, 6_000_000, 800_000)


def _ordinals(prs, bad: bool) -> None:
    slide = blank(prs)
    add_text(slide, "2-го числа" if bad else "второго числа", 400_000, 400_000, 6_000_000, 800_000)


def _list(prs, bad: bool) -> None:
    slide = blank(prs)
    box = add_text(slide, "", 400_000, 400_000, 6_000_000, 1_500_000)
    frame = box.text_frame
    frame.clear()
    first = frame.paragraphs[0]
    first.text = "Первый."
    set_bullet(first)
    second = frame.add_paragraph()
    second.text = "Второй" if bad else "Второй."
    set_bullet(second)


def _heuristic(prs, bad: bool) -> None:
    slide = blank(prs)
    if bad:
        add_text(slide, "Низко", 400_000, 4_000_000, 6_000_000, 800_000, size=32)
    else:
        add_text(slide, "Заголовок", 400_000, 200_000, 6_000_000, 800_000, size=32)


def _font(prs, bad: bool) -> None:
    slide = blank(prs)
    add_text(slide, "Размер", 400_000, 400_000, 4_000_000, 800_000, size=14 if bad else 18)


def _margin(prs, bad: bool) -> None:
    slide = blank(prs)
    if bad:
        add_text(slide, "Низ", 400_000, SLIDE_H - 500_000, 4_000_000, 480_000)
    else:
        add_text(slide, "Середина", 400_000, 1_000_000, 4_000_000, 800_000)


def _numbers(prs, bad: bool) -> None:
    slide = blank(prs)
    add_text(slide, "100000 человек" if bad else "9999 человек", 400_000, 400_000, 6_000_000, 800_000)


def _english(prs, bad: bool) -> None:
    slide = blank(prs)
    add_text(slide, "Hello world" if bad else "Привет мир", 400_000, 400_000, 6_000_000, 800_000)


def _units(prs, bad: bool) -> None:
    slide = blank(prs)
    box = add_text(slide, "", 400_000, 400_000, 4_000_000, 800_000)
    paragraph = box.text_frame.paragraphs[0]
    paragraph.clear()
    first = paragraph.add_run()
    first.text = "10 м2" if bad else "10 м"
    first.font.size = Pt(24)
    first.font.name = "Calibri"
    if not bad:
        second = paragraph.add_run()
        second.text = "2"
        second.font.size = Pt(24)
        second.font.name = "Calibri"
        second.font._element.set("baseline", "30000")


def _overlap(prs, bad: bool) -> None:
    slide = blank(prs)
    add_text(slide, "Первый", 400_000, 400_000, 3_000_000, 1_000_000)
    if bad:
        add_text(slide, "Второй", 500_000, 500_000, 3_000_000, 1_000_000)
    else:
        add_text(slide, "Второй", 4_000_000, 2_000_000, 3_000_000, 1_000_000)


def _contrast(prs, bad: bool) -> None:
    slide = blank(prs)
    box = add_text(slide, "Контраст", 400_000, 400_000, 4_000_000, 800_000, size=16)
    run = box.text_frame.paragraphs[0].runs[-1]
    box.fill.solid()
    if bad:
        run.font.color.rgb = RGBColor(255, 255, 255)
        box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    else:
        run.font.color.rgb = RGBColor(0, 0, 0)
        box.fill.fore_color.rgb = RGBColor(255, 255, 255)


def _hanging(prs, bad: bool) -> None:
    slide = blank(prs)
    box = add_text(slide, "", 400_000, 400_000, 6_000_000, 1_200_000)
    paragraph = box.text_frame.paragraphs[0]
    paragraph.clear()
    first = paragraph.add_run()
    first.text = "Текст в" if bad else "Текст внутри"
    first.font.size = Pt(24)
    first.font.name = "Calibri"
    paragraph.add_line_break()
    second = paragraph.add_run()
    second.text = "начале"
    second.font.size = Pt(24)
    second.font.name = "Calibri"


CASES = [
    ("aspect_16_9", _aspect),
    ("embedded_fonts", _fonts),
    ("slide_duplicates", _duplicates),
    ("missing_title", _missing_title),
    ("title_dot", _title_dot),
    ("has_video", _video),
    ("quote_guillemets", _quotes),
    ("em_dash", _dash),
    ("ordinals", _ordinals),
    ("list_periods", _list),
    ("title_heuristic", _heuristic),
    ("font_size_min", _font),
    ("bottom_margin", _margin),
    ("number_grouping", _numbers),
    ("english_ratio", _english),
    ("superscript_units", _units),
    ("overlap", _overlap),
    ("contrast_fill", _contrast),
    ("hanging_prepositions", _hanging),
]


def test_each_check_flags_only_the_broken_slide():
    for check_id, builder in CASES:
        assert_clean(builder, check_id)
        assert_flagged(builder, check_id)


def test_duplicate_message_lists_slide_numbers():
    issue = assert_flagged(_duplicates, "slide_duplicates")
    assert issue["slide_numbers"] == [1, 2]
    assert issue["message"] == "Слайды 1 и 2 совпадают по тексту и набору объектов."


def test_font_size_message_and_footer_ignored():
    issue = assert_flagged(_font, "font_size_min")
    assert issue["message"] == "Размер шрифта 14 pt меньше минимального (18 pt)."
    prs = presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.is_placeholder and shape.placeholder_format.type is not None:
            name = str(shape.placeholder_format.type).lower()
            if "footer" in name or "date" in name or "slide_number" in name:
                shape.text_frame.paragraphs[0].clear()
                run = shape.text_frame.paragraphs[0].add_run()
                run.text = "мелкий"
                run.font.size = Pt(10)
    slide.shapes.title.text = "Тема"
    assert issues_for(save(prs), "font_size_min") == []


def test_english_in_monospace_is_ignored():
    prs = presentation()
    slide = blank(prs)
    add_text(slide, "HelloWorldCode", 400_000, 400_000, 6_000_000, 600_000, font="Consolas")
    add_text(slide, "Обычный русский текст", 400_000, 1_200_000, 6_000_000, 600_000, font="Calibri")
    assert issues_for(save(prs), "english_ratio") == []


def test_numeric_range_uses_em_dash_rule():
    prs = presentation()
    slide = blank(prs)
    add_text(slide, "годы 1990-2000", 400_000, 400_000, 6_000_000, 800_000)
    found = issues_for(save(prs), "em_dash")
    assert len(found) == 1


def test_grouped_number():
    prs = presentation()
    slide = blank(prs)
    add_text(slide, "Бюджет 12 345 рублей", 400_000, 400_000, 6_000_000, 800_000)
    found = issues_for(save(prs), "number_grouping")
    assert len(found) == 1
    assert "тыс." in found[0]["message"]


def test_contrast_skips_gradient_and_transparent_fill():
    prs = presentation()
    slide = blank(prs)
    box = add_text(slide, "Текст", 400_000, 400_000, 4_000_000, 800_000, size=16)
    run = box.text_frame.paragraphs[0].runs[-1]
    run.font.color.rgb = RGBColor(255, 255, 255)
    box.fill.gradient()
    assert issues_for(save(prs), "contrast_fill") == []

    prs = presentation()
    slide = blank(prs)
    box = add_text(slide, "Текст", 400_000, 400_000, 4_000_000, 800_000, size=16)
    run = box.text_frame.paragraphs[0].runs[-1]
    run.font.color.rgb = RGBColor(255, 255, 255)
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(255, 255, 255)
    solid = box._element.spPr.find(qn("a:solidFill"))
    color = solid.find(qn("a:srgbClr"))
    alpha = etree.SubElement(color, qn("a:alpha"))
    alpha.set("val", "50000")
    assert issues_for(save(prs), "contrast_fill") == []

