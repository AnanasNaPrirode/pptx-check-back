from __future__ import annotations

import base64
from io import BytesIO

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

SLIDE_W = 12_192_000
SLIDE_H = 6_858_000

MIN_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    return prs


def save(prs: Presentation) -> bytes:
    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_text(slide, text: str, x: int, y: int, w: int, h: int, size: float = 24, font: str = "Calibri"):
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    paragraph = box.text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.name = font
    return box


def set_bullet(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bullet = etree.SubElement(p_pr, qn("a:buChar"))
    bullet.set("char", "•")
